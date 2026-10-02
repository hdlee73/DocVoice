"""받아쓰기 결과를 xlsx / docx / pdf / txt 로 저장."""
from __future__ import annotations

from pathlib import Path
from xml.sax.saxutils import escape

from .segment import Paragraph, Sentence, fmt_time, group_paragraphs

FORMATS = ("xlsx", "docx", "pdf", "txt")
ASSETS = Path(__file__).parent / "assets"

# 파스텔 팔레트 (앱 UI와 동일 계열)
LAVENDER = "E9E4FA"
MINT = "E3F4EC"
ROW_ALT = "F7F5FD"
INK = "4A4560"


def _speaker_label(spk: int | None) -> str:
    return f"화자 {spk + 1}" if spk is not None else ""


# ---------------------------------------------------------------------------
# XLSX : 문장 1개 = 1행
# ---------------------------------------------------------------------------

def write_xlsx(sentences: list[Sentence], path: Path, include_time: bool = True,
               show_speaker: bool = False) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "문장"
    headers = ["번호"]
    if include_time:
        headers += ["시작", "종료"]
    if show_speaker:
        headers.append("화자")
    headers.append("문장")
    ws.append(headers)

    head_fill = PatternFill("solid", fgColor=LAVENDER)
    alt_fill = PatternFill("solid", fgColor=ROW_ALT)
    thin = Side(style="thin", color="E4E0F2")
    border = Border(bottom=thin)
    for c in ws[1]:
        c.font = Font(bold=True, color=INK, name="Malgun Gothic")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border

    for i, s in enumerate(sentences, 1):
        row = [i]
        if include_time:
            row += [fmt_time(s.start), fmt_time(s.end)]
        if show_speaker:
            row.append(_speaker_label(s.speaker))
        row.append(s.text)
        ws.append(row)
        r = ws.max_row
        for c in ws[r]:
            c.font = Font(color=INK, name="Malgun Gothic")
            c.alignment = Alignment(vertical="top", wrap_text=(c.column == len(row)))
            c.border = border
            if i % 2 == 0:
                c.fill = alt_fill
        text_cell = ws.cell(row=r, column=len(row))
        text_cell.data_type = "s"  # '=' 로 시작해도 수식으로 해석하지 않음

    widths = {"번호": 7, "시작": 12, "종료": 12, "화자": 9, "문장": 90}
    for idx, h in enumerate(headers, 1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = widths[h]
    ws.freeze_panes = "A2"
    wb.save(str(path))
    return path


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def write_docx(paragraphs: list[Paragraph], path: Path, title: str = "",
               include_time: bool = True, show_speaker: bool = False) -> Path:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Malgun Gothic"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.5

    if title:
        h = doc.add_paragraph()
        r = h.add_run(title)
        r.bold = True
        r.font.size = Pt(16)
        r.font.color.rgb = RGBColor(0x6B, 0x5B, 0xB5)
    for para in paragraphs:
        p = doc.add_paragraph()
        prefix = []
        if include_time:
            prefix.append(f"[{fmt_time(para.start)[:-2]}]")
        if show_speaker and para.speaker is not None:
            prefix.append(_speaker_label(para.speaker) + ":")
        if prefix:
            tag = p.add_run(" ".join(prefix) + " ")
            tag.font.size = Pt(9)
            tag.font.color.rgb = RGBColor(0x9A, 0x90, 0xC0)
        p.add_run(para.text)
    doc.save(str(path))
    return path


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

_FONT_CANDIDATES = [
    (ASSETS / "NanumGothic-Regular.ttf", 0),
    (Path(r"C:\Windows\Fonts\malgun.ttf"), 0),
    (Path("/System/Library/Fonts/AppleSDGothicNeo.ttc"), 0),
    (Path("/System/Library/Fonts/Supplemental/AppleGothic.ttf"), 0),
    (Path("/usr/share/fonts/truetype/nanum/NanumGothic.ttf"), 0),
]
_registered: str | None = None


def _register_korean_font() -> str:
    global _registered
    if _registered:
        return _registered
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for path, idx in _FONT_CANDIDATES:
        if path.exists():
            try:
                pdfmetrics.registerFont(TTFont("DocVoiceKR", str(path), subfontIndex=idx))
                _registered = "DocVoiceKR"
                return _registered
            except Exception:  # noqa: BLE001
                continue
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont

    pdfmetrics.registerFont(UnicodeCIDFont("HYSMyeongJo-Medium"))
    _registered = "HYSMyeongJo-Medium"
    return _registered


def write_pdf(paragraphs: list[Paragraph], path: Path, title: str = "",
              include_time: bool = True, show_speaker: bool = False) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph as RLPara
    from reportlab.platypus import SimpleDocTemplate, Spacer

    font = _register_korean_font()
    body = ParagraphStyle("body", fontName=font, fontSize=10.5, leading=17,
                          textColor=colors.HexColor("#" + INK), spaceAfter=8)
    head = ParagraphStyle("head", parent=body, fontSize=16, leading=22,
                          textColor=colors.HexColor("#6B5BB5"), spaceAfter=14)
    tag = ParagraphStyle("tag", parent=body, fontSize=8.5, leading=12,
                         textColor=colors.HexColor("#9A90C0"), spaceAfter=1)

    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=22 * mm, rightMargin=22 * mm,
                            topMargin=22 * mm, bottomMargin=22 * mm, title=title or path.stem)
    flow = []
    if title:
        flow.append(RLPara(escape(title), head))
    for para in paragraphs:
        label = []
        if include_time:
            label.append(f"[{fmt_time(para.start)[:-2]}]")
        if show_speaker and para.speaker is not None:
            label.append(_speaker_label(para.speaker))
        if label:
            flow.append(RLPara(escape(" ".join(label)), tag))
        flow.append(RLPara(escape(para.text), body))
    if not flow:
        flow.append(Spacer(1, 1))
    doc.build(flow)
    return path


def write_txt(paragraphs: list[Paragraph], path: Path, include_time: bool = True,
              show_speaker: bool = False) -> Path:
    blocks = []
    for para in paragraphs:
        head = []
        if include_time:
            head.append(f"[{fmt_time(para.start)[:-2]}]")
        if show_speaker and para.speaker is not None:
            head.append(_speaker_label(para.speaker) + ":")
        blocks.append((" ".join(head) + " " if head else "") + para.text)
    path.write_text("\n\n".join(blocks) + "\n", encoding="utf-8")
    return path


def _unique(path: Path) -> Path:
    if not path.exists():
        return path
    k = 1
    while True:
        cand = path.with_name(f"{path.stem}_{k}{path.suffix}")
        if not cand.exists():
            return cand
        k += 1


def export_all(
    sentences: list[Sentence],
    out_dir: Path,
    stem: str,
    formats: list[str],
    gap: float = 1.5,
    use_speaker: bool = True,
    include_time: bool = True,
    show_speaker: bool = False,
) -> list[Path]:
    paragraphs = group_paragraphs(sentences, gap=gap, use_speaker=use_speaker)
    out: list[Path] = []
    for fmt in formats:
        p = _unique(Path(out_dir) / f"{stem}.{fmt}")
        if fmt == "xlsx":
            write_xlsx(sentences, p, include_time, show_speaker)
        elif fmt == "docx":
            write_docx(paragraphs, p, stem, include_time, show_speaker)
        elif fmt == "pdf":
            write_pdf(paragraphs, p, stem, include_time, show_speaker)
        elif fmt == "txt":
            write_txt(paragraphs, p, include_time, show_speaker)
        out.append(p)
    return out
