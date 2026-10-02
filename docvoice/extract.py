"""문서 → 텍스트 추출 (txt, md, csv, pdf, doc/docx, xls/xlsx, ppt/pptx, hwp/hwpx)."""
from __future__ import annotations

import csv
import io
import os
import re
import shutil
import struct
import subprocess
import tempfile
import zipfile
import zlib
from pathlib import Path
from typing import Callable

SUPPORTED_EXTS = (
    ".pdf", ".docx", ".doc", ".xlsx", ".xlsm", ".xls", ".pptx", ".ppt",
    ".hwpx", ".hwp", ".txt", ".md", ".csv",
)


class ExtractError(Exception):
    pass


# ---------------------------------------------------------------------------
# 공통 유틸
# ---------------------------------------------------------------------------

def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\u00a0", " ")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    text = re.sub(r"[\ue000-\uf8ff]", "", text)  # 사용자 정의 영역(HWP 옛한글 잔여 등)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _fmt_cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _row_line(cells) -> str:
    vals = []
    for c in cells:
        s = _fmt_cell(c)
        if s and (not vals or vals[-1] != s):  # 병합 셀 중복 제거
            vals.append(s)
    return ", ".join(vals)


def _reflow(text: str) -> str:
    """PDF처럼 줄 중간에서 끊긴 문장을 이어 붙인다."""
    lines = text.split("\n")
    out: list[str] = []
    for ln in lines:
        s = ln.strip()
        if not s:
            out.append("")
            continue
        if out and out[-1]:
            prev = out[-1]
            bullet = re.match(r"^([\-•·▪●○■□◆◇▶※*]|\d+[.)]|[가-하][.)])\s", s)
            if prev[-1] not in ".?!。！？…:;" and not bullet:
                out[-1] = prev + " " + s
                continue
        out.append(s)
    return "\n".join(out)


def _read_text_file(p: Path) -> str:
    raw = p.read_bytes()
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return raw.decode("utf-16", errors="ignore")
    for enc in ("utf-8-sig", "cp949", "euc-kr"):
        try:
            return raw.decode(enc)
        except UnicodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


# ---------------------------------------------------------------------------
# 각 형식
# ---------------------------------------------------------------------------

def _from_csv(p: Path) -> str:
    text = _read_text_file(p)
    rows = csv.reader(io.StringIO(text))
    return "\n".join(_row_line(r) for r in rows)


def _from_pdf(p: Path) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(p))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ExtractError("암호가 걸린 PDF입니다.")
        pages = [(pg.extract_text() or "") for pg in reader.pages]
    except ExtractError:
        raise
    except Exception as e:  # noqa: BLE001
        raise ExtractError(f"PDF를 읽을 수 없습니다: {e}") from e
    text = "\n\n".join(pages)
    if not text.strip():
        raise ExtractError("PDF에서 텍스트를 찾지 못했습니다. (스캔본 PDF는 OCR이 필요합니다)")
    return _reflow(text)


def _from_docx(p: Path) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        doc = Document(str(p))
    except Exception as e:  # noqa: BLE001
        raise ExtractError(f"DOCX를 읽을 수 없습니다: {e}") from e
    lines: list[str] = []
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            lines.append(Paragraph(child, doc).text)
        elif child.tag.endswith("}tbl"):
            for row in Table(child, doc).rows:
                lines.append(_row_line(c.text for c in row.cells))
    return "\n".join(lines)


def _from_xlsx(p: Path) -> str:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(str(p), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise ExtractError(f"XLSX를 읽을 수 없습니다: {e}") from e
    parts: list[str] = []
    for ws in wb.worksheets:
        lines = [_row_line(r) for r in ws.iter_rows(values_only=True)]
        lines = [ln for ln in lines if ln]
        if lines:
            parts.append("\n".join(lines))
    wb.close()
    return "\n\n".join(parts)


def _from_xls(p: Path) -> str:
    import xlrd

    try:
        wb = xlrd.open_workbook(str(p))
    except Exception as e:  # noqa: BLE001
        raise ExtractError(f"XLS를 읽을 수 없습니다: {e}") from e
    parts: list[str] = []
    for sh in wb.sheets():
        lines = [_row_line(sh.row_values(r)) for r in range(sh.nrows)]
        lines = [ln for ln in lines if ln]
        if lines:
            parts.append("\n".join(lines))
    return "\n\n".join(parts)


def _shape_texts(shape) -> list[str]:
    out: list[str] = []
    if getattr(shape, "shapes", None) is not None and shape.shape_type == 6:  # GROUP
        for sub in shape.shapes:
            out.extend(_shape_texts(sub))
        return out
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            out.append(_row_line(c.text for c in row.cells))
    elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            t = "".join(r.text for r in para.runs) or para.text
            out.append(t)
    return out


def _from_pptx(p: Path) -> str:
    from pptx import Presentation

    try:
        prs = Presentation(str(p))
    except Exception as e:  # noqa: BLE001
        raise ExtractError(f"PPTX를 읽을 수 없습니다: {e}") from e
    slides: list[str] = []
    for slide in prs.slides:
        lines: list[str] = []
        for shape in slide.shapes:
            lines.extend(_shape_texts(shape))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            note = slide.notes_slide.notes_text_frame.text.strip()
            if note:
                lines.append(note)
        if any(x.strip() for x in lines):
            slides.append("\n".join(lines))
    return "\n\n".join(slides)


# ---- HWPX ----------------------------------------------------------------

def _from_hwpx(p: Path) -> str:
    import xml.etree.ElementTree as ET

    try:
        zf = zipfile.ZipFile(str(p))
    except zipfile.BadZipFile as e:
        raise ExtractError("HWPX 파일이 손상되었거나 형식이 올바르지 않습니다.") from e
    names = [n for n in zf.namelist() if re.match(r"Contents/section\d+\.xml$", n, re.I)]
    if not names:
        raise ExtractError("HWPX에서 본문(section) 데이터를 찾지 못했습니다. (암호 문서일 수 있습니다)")
    names.sort(key=lambda n: int(re.search(r"(\d+)\.xml", n).group(1)))
    lines: list[str] = []
    for name in names:
        buf: list[str] = []
        for _ev, el in ET.iterparse(io.BytesIO(zf.read(name)), events=("end",)):
            local = el.tag.rsplit("}", 1)[-1]
            if local == "t":
                buf.append("".join(el.itertext()))
            elif local == "p":
                lines.append("".join(buf))
                buf = []
        if buf:
            lines.append("".join(buf))
    return "\n".join(lines)


# ---- HWP (5.x 바이너리) ---------------------------------------------------

_HWPTAG_PARA_TEXT = 67


def _hwp_para_text(rec: bytes) -> str:
    buf = bytearray()
    i, n = 0, len(rec)
    while i + 2 <= n:
        code = rec[i] | (rec[i + 1] << 8)
        i += 2
        if code >= 32:
            buf += rec[i - 2:i]
        elif code in (10, 13):
            buf += b"\n\x00"
        elif code in (30, 31):  # 묶음/고정폭 빈칸
            buf += b" \x00"
        elif code in (0, 24, 25, 26, 27, 28, 29):
            continue
        else:  # 인라인/확장 컨트롤: 코드 포함 8 WCHAR → 나머지 14바이트 건너뜀
            if code == 9:
                buf += b"\t\x00"
            i += 14
    return buf.decode("utf-16-le", errors="ignore")


def _parse_hwp_section(data: bytes) -> str:
    pos, n = 0, len(data)
    paras: list[str] = []
    while pos + 4 <= n:
        (h,) = struct.unpack_from("<I", data, pos)
        pos += 4
        tag = h & 0x3FF
        size = (h >> 20) & 0xFFF
        if size == 0xFFF:
            if pos + 4 > n:
                break
            (size,) = struct.unpack_from("<I", data, pos)
            pos += 4
        rec = data[pos:pos + size]
        pos += size
        if tag == _HWPTAG_PARA_TEXT:
            paras.append(_hwp_para_text(rec))
    paras = [t.rstrip("\n") for t in paras]
    return "\n".join(t for t in paras if t.strip())


def _from_hwp(p: Path) -> str:
    import olefile

    if not olefile.isOleFile(str(p)):
        raise ExtractError("지원하지 않는 HWP 형식입니다. (HWP 5.x 또는 HWPX만 지원)")
    ole = olefile.OleFileIO(str(p))
    try:
        header = ole.openstream("FileHeader").read()
        if not header.startswith(b"HWP Document File"):
            raise ExtractError("HWP 헤더가 올바르지 않습니다.")
        (flags,) = struct.unpack("<I", header[36:40])
        if flags & 0x02:
            raise ExtractError("암호가 걸린 HWP 문서는 읽을 수 없습니다.")
        if flags & 0x04:
            raise ExtractError("배포용 HWP 문서는 읽을 수 없습니다. 한글에서 일반 문서로 저장 후 다시 시도하세요.")
        compressed = bool(flags & 0x01)
        sections = sorted(
            (d for d in ole.listdir() if len(d) == 2 and d[0] == "BodyText" and d[1].startswith("Section")),
            key=lambda d: int(re.sub(r"\D", "", d[1]) or 0),
        )
        if not sections:
            raise ExtractError("HWP에서 본문을 찾지 못했습니다.")
        out: list[str] = []
        for d in sections:
            data = ole.openstream(d).read()
            if compressed:
                data = zlib.decompress(data, -15)
            out.append(_parse_hwp_section(data))
        return "\n".join(out)
    finally:
        ole.close()


# ---- 레거시 DOC / PPT : LibreOffice 우선, 없으면 내장 파서 ------------------

def _find_soffice() -> str | None:
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    for cand in (
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ):
        if os.path.exists(cand):
            return cand
    return None


def _convert_with_soffice(p: Path, fmt: str) -> Path | None:
    exe = _find_soffice()
    if not exe:
        return None
    tmp = Path(tempfile.mkdtemp(prefix="docvoice_"))
    try:
        subprocess.run(
            [exe, "--headless", "--convert-to", fmt, "--outdir", str(tmp), str(p)],
            capture_output=True, timeout=240, check=True,
        )
    except Exception:  # noqa: BLE001
        shutil.rmtree(tmp, ignore_errors=True)
        return None
    out = tmp / (p.stem + "." + fmt)
    return out if out.exists() else None


def _legacy_via_soffice(p: Path, fmt: str, reader: Callable[[Path], str]) -> str | None:
    conv = _convert_with_soffice(p, fmt)
    if conv is None:
        return None
    try:
        return reader(conv)
    finally:
        shutil.rmtree(conv.parent, ignore_errors=True)


def _parse_ppt_records(data: bytes, out: list[str], in_slide: bool = False) -> None:
    pos, n = 0, len(data)
    while pos + 8 <= n:
        verinst, rtype, rlen = struct.unpack_from("<HHI", data, pos)
        pos += 8
        end = min(pos + rlen, n)
        ver = verinst & 0xF
        if ver == 0xF:  # 컨테이너
            slide_ctx = in_slide or rtype == 1006 or (rtype == 4080 and (verinst >> 4) == 0)
            if rtype in (1016, 1008):  # 마스터 / 노트 제외
                pass
            else:
                _parse_ppt_records(data[pos:end], out, slide_ctx)
        elif in_slide and rtype == 0x0FA0:
            out.append(data[pos:end].decode("utf-16-le", errors="ignore"))
        elif in_slide and rtype == 0x0FA8:
            out.append(data[pos:end].decode("cp1252", errors="ignore"))
        pos = end


def _from_ppt_builtin(p: Path) -> str:
    import olefile

    if not olefile.isOleFile(str(p)):
        raise ExtractError("PPT 파일 형식이 올바르지 않습니다.")
    ole = olefile.OleFileIO(str(p))
    try:
        data = ole.openstream("PowerPoint Document").read()
    finally:
        ole.close()
    texts: list[str] = []
    _parse_ppt_records(data, texts)
    seen, uniq = set(), []
    for t in texts:
        t = t.replace("\r", "\n").replace("\x0b", "\n").strip()
        if t and t not in seen:
            seen.add(t)
            uniq.append(t)
    if not uniq:
        raise ExtractError("PPT에서 텍스트를 찾지 못했습니다.")
    return "\n".join(uniq)


def _from_ppt(p: Path) -> str:
    text = _legacy_via_soffice(p, "pptx", _from_pptx)
    if text is not None:
        return text
    return _from_ppt_builtin(p)


def _doc_text_from_stream(word: bytes, table: bytes) -> str:
    (ccp_text,) = struct.unpack_from("<i", word, 0x4C)
    fc_clx, lcb_clx = struct.unpack_from("<II", word, 0x1A2)
    clx = table[fc_clx:fc_clx + lcb_clx]
    i = 0
    while i < len(clx) and clx[i] == 1:  # Prc 건너뜀
        (cb,) = struct.unpack_from("<H", clx, i + 1)
        i += 3 + cb
    if i >= len(clx) or clx[i] != 2:
        raise ExtractError("DOC 구조를 해석하지 못했습니다.")
    (lcb,) = struct.unpack_from("<I", clx, i + 1)
    plc = clx[i + 5:i + 5 + lcb]
    n = (lcb - 4) // 12
    cps = struct.unpack_from("<%di" % (n + 1), plc, 0)
    parts: list[str] = []
    for k in range(n):
        (fc,) = struct.unpack_from("<I", plc, (n + 1) * 4 + k * 8 + 2)
        count = cps[k + 1] - cps[k]
        if fc & 0x40000000:
            off = (fc & 0x3FFFFFFF) // 2
            parts.append(word[off:off + count].decode("cp1252", errors="ignore"))
        else:
            parts.append(word[fc:fc + count * 2].decode("utf-16-le", errors="ignore"))
    text = "".join(parts)[:ccp_text]
    text = re.sub(r"\x13[^\x13\x14\x15]*\x14", "", text)  # 필드 코드 제거
    text = text.replace("\x07", " ").replace("\x0b", "\n").replace("\x0c", "\n")
    text = text.replace("\r", "\n")
    return re.sub(r"[\x00-\x08\x0e-\x1f]", "", text)


def _from_doc_builtin(p: Path) -> str:
    import olefile

    if not olefile.isOleFile(str(p)):
        raise ExtractError("DOC 파일 형식이 올바르지 않습니다.")
    ole = olefile.OleFileIO(str(p))
    try:
        word = ole.openstream("WordDocument").read()
        (flags,) = struct.unpack_from("<H", word, 0x0A)
        if flags & 0x0100:
            raise ExtractError("암호가 걸린 DOC 문서는 읽을 수 없습니다.")
        tname = "1Table" if flags & 0x0200 else "0Table"
        table = ole.openstream(tname).read()
    finally:
        ole.close()
    try:
        return _doc_text_from_stream(word, table)
    except ExtractError:
        raise
    except Exception as e:  # noqa: BLE001
        raise ExtractError(f"DOC를 해석하지 못했습니다: {e}") from e


def _from_doc(p: Path) -> str:
    text = _legacy_via_soffice(p, "docx", _from_docx)
    if text is not None:
        return text
    return _from_doc_builtin(p)


_HANDLERS: dict[str, Callable[[Path], str]] = {
    ".txt": _read_text_file,
    ".md": _read_text_file,
    ".csv": _from_csv,
    ".pdf": _from_pdf,
    ".docx": _from_docx,
    ".doc": _from_doc,
    ".xlsx": _from_xlsx,
    ".xlsm": _from_xlsx,
    ".xls": _from_xls,
    ".pptx": _from_pptx,
    ".ppt": _from_ppt,
    ".hwpx": _from_hwpx,
    ".hwp": _from_hwp,
}


def extract_text(path: str | os.PathLike) -> str:
    p = Path(path)
    if not p.exists():
        raise ExtractError(f"파일을 찾을 수 없습니다: {p}")
    handler = _HANDLERS.get(p.suffix.lower())
    if handler is None:
        raise ExtractError(f"지원하지 않는 형식입니다: {p.suffix} (지원: {', '.join(SUPPORTED_EXTS)})")
    text = clean_text(handler(p))
    if not text:
        raise ExtractError("문서에서 읽을 텍스트를 찾지 못했습니다.")
    return text
