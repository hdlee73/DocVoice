
from docvoice import tts
from docvoice.export import export_all
from docvoice.segment import Sentence, assign_speakers


def _sents():
    s = [Sentence(0, 2, "안녕하세요, 회의를 시작하겠습니다."), Sentence(2.2, 4, "Thank you all for coming."),
         Sentence(4.1, 7, "=SUM(A1) 로 시작하는 문장도 수식이 아니어야 합니다."),
         Sentence(12, 14, "A line after a long pause, indeed.")]
    assign_speakers(s, [(0, 4.05, 0), (4.05, 14, 1)])
    return s


def test_export_all_formats(tmp_path):
    from docx import Document
    from openpyxl import load_workbook
    from pypdf import PdfReader

    paths = export_all(_sents(), tmp_path, "회의", ["xlsx", "docx", "pdf", "txt"], gap=1.5)
    assert [p.suffix for p in paths] == [".xlsx", ".docx", ".pdf", ".txt"]

    ws = load_workbook(paths[0]).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("번호", "시작", "종료", "문장")
    assert len(rows) == 5  # 헤더 + 문장 4개 → 한 문장 한 행
    assert rows[3][3].startswith("=SUM")
    assert ws.cell(row=4, column=4).data_type == "s"

    paras = [p.text for p in Document(paths[1]).paragraphs if p.text.strip()]
    assert len(paras) == 1 + 3  # 제목 + (화자1) + (화자2) + (긴 침묵 뒤)

    text = "".join(pg.extract_text() for pg in PdfReader(paths[2]).pages)
    assert "Thank you all" in text and "안녕하세요" in text

    blocks = paths[3].read_text(encoding="utf-8").strip().split("\n\n")
    assert len(blocks) == 3


def test_export_without_time_and_unique_names(tmp_path):
    export_all(_sents(), tmp_path, "a", ["xlsx"], include_time=False)
    p2 = export_all(_sents(), tmp_path, "a", ["xlsx"], include_time=False)[0]
    assert p2.name == "a_1.xlsx"
    from openpyxl import load_workbook

    assert [c.value for c in load_workbook(p2).active[1]] == ["번호", "문장"]


def test_build_chunks_language_split():
    chunks = tts.build_chunks("안녕하세요. 반갑습니다.\nHello there. How are you?\nAI 모델을 사용합니다.")
    assert [c.lang for c in chunks] == ["ko", "en", "ko"]
    assert chunks[1].text.startswith("Hello there.")


def test_build_chunks_size_limit():
    text = "\n".join(f"Sentence number {i} is here." for i in range(400))
    chunks = tts.build_chunks(text, max_chars=500)
    assert len(chunks) > 5 and all(len(c.text) <= 520 for c in chunks)


def test_convert_document_to_mp3_with_fake_synth(tmp_path, monkeypatch):
    calls = []

    async def fake(text, voice, rate, fallback, retries=3):
        calls.append((voice, rate))
        return voice.encode()

    monkeypatch.setattr(tts, "_synth_one", fake)
    src = tmp_path / "doc.txt"
    src.write_text("안녕하세요. 반갑습니다.\nGood morning everyone.", encoding="utf-8")
    out = tts.convert_document_to_mp3(src, tmp_path / "out", accent="uk", speed=1.2)
    assert out.read_bytes() == b"ko-KR-InJoonNeuralen-GB-RyanNeural"
    assert {c[1] for c in calls} == {"+20%"}


def test_convert_audio_to_documents_pipeline(tmp_path, monkeypatch):
    import numpy as np

    from docvoice import diarize as dz
    from docvoice import stt
    from docvoice.segment import Segment

    segs = [Segment(0, 3, "Hello everyone, welcome to the show."), Segment(3.1, 4, "Yes."),
            Segment(4.2, 8, "안녕하세요. 오늘 주제는 인공지능입니다."), Segment(15, 18, "마무리하겠습니다.")]
    monkeypatch.setattr(stt, "transcribe", lambda *a, **k: (segs, np.zeros(16000 * 20, np.float32), "en"))
    monkeypatch.setattr(dz, "diarize", lambda *a, **k: [(0, 4.1, 0), (4.1, 20, 1)])
    audio = tmp_path / "talk.wav"
    audio.write_bytes(b"x")
    outs = stt.convert_audio_to_documents(audio, tmp_path / "o", ["xlsx", "docx"], gap=1.5)
    from openpyxl import load_workbook

    rows = list(load_workbook(outs[0]).active.iter_rows(values_only=True))[1:]
    assert [r[3] for r in rows] == ["Hello everyone, welcome to the show. Yes.", "안녕하세요.",
                                    "오늘 주제는 인공지능입니다.", "마무리하겠습니다."]
