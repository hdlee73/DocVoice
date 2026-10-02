import struct
import subprocess
import zipfile
from pathlib import Path

import pytest

from docvoice import extract as ex

HAS_SOFFICE = ex._find_soffice() is not None


def test_txt_encodings(tmp_path):
    p = tmp_path / "a.txt"
    p.write_bytes("안녕하세요\nhello".encode("cp949"))
    assert ex.extract_text(p) == "안녕하세요\nhello"


def test_docx_with_table(tmp_path):
    from docx import Document

    d = Document()
    d.add_paragraph("첫 문단입니다.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "이름", "점수"
    t.cell(1, 0).text, t.cell(1, 1).text = "철수", "90"
    d.add_paragraph("마지막 문단.")
    p = tmp_path / "a.docx"
    d.save(p)
    assert ex.extract_text(p) == "첫 문단입니다.\n이름, 점수\n철수, 90\n마지막 문단."


def test_xlsx_and_csv(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["제품", "수량"])
    ws.append(["사과", 3.0])
    ws2 = wb.create_sheet("s2")
    ws2.append(["Hello", "World"])
    p = tmp_path / "a.xlsx"
    wb.save(p)
    assert ex.extract_text(p) == "제품, 수량\n사과, 3\n\nHello, World"
    c = tmp_path / "a.csv"
    c.write_text("a,b\n1,2\n", encoding="utf-8")
    assert ex.extract_text(c) == "a, b\n1, 2"


def test_pptx(tmp_path):
    from pptx import Presentation

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "제목 슬라이드"
    s.placeholders[1].text = "본문 내용"
    p = tmp_path / "a.pptx"
    prs.save(p)
    assert ex.extract_text(p) == "제목 슬라이드\n본문 내용"


def test_pdf(tmp_path):
    from docvoice.export import write_pdf
    from docvoice.segment import Paragraph, Sentence

    p = tmp_path / "a.pdf"
    write_pdf([Paragraph([Sentence(0, 1, "Hello PDF world.")])], p, include_time=False)
    assert "Hello PDF world." in ex.extract_text(p)


def test_hwpx(tmp_path):
    ns = 'xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph"'
    sec = (f'<?xml version="1.0" encoding="UTF-8"?><hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" {ns}>'
           '<hp:p><hp:run><hp:t>안녕하세요 </hp:t><hp:t>한글 문서입니다.</hp:t></hp:run></hp:p>'
           '<hp:p><hp:run><hp:t>둘째 문단</hp:t></hp:run></hp:p></hs:sec>')
    p = tmp_path / "a.hwpx"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("mimetype", "application/hwp+zip")
        z.writestr("Contents/section0.xml", sec)
    assert ex.extract_text(p) == "안녕하세요 한글 문서입니다.\n둘째 문단"


def _hwp_record(tag, payload):
    size = len(payload)
    if size >= 0xFFF:
        return struct.pack("<II", tag | (0xFFF << 20), size) + payload
    return struct.pack("<I", tag | (size << 20)) + payload


def test_hwp_section_parser():
    text = "안녕 HWP".encode("utf-16-le")
    ctrl = struct.pack("<H", 2) + b"\x00" * 14  # 확장 컨트롤(구역 정의 등) → 건너뛰어야 함
    para1 = ctrl + text + struct.pack("<H", 13)
    para2 = "둘째 문단😀".encode("utf-16-le") + struct.pack("<H", 13)
    data = (_hwp_record(66, b"\x00" * 8) + _hwp_record(67, para1) + _hwp_record(67, para2)
            + _hwp_record(67, "긴".encode("utf-16-le") * 3000))
    out = ex._parse_hwp_section(data)
    lines = out.split("\n")
    assert lines[0] == "안녕 HWP" and lines[1] == "둘째 문단😀"
    assert lines[2] == "긴" * 3000


def test_ppt_record_parser_skips_masters():
    def rec(ver, inst, rtype, body):
        return struct.pack("<HHI", (inst << 4) | ver, rtype, len(body)) + body

    master = rec(0xF, 0, 1016, rec(0, 0, 0x0FA0, "마스터 서식".encode("utf-16-le")))
    slide = rec(0xF, 0, 1006, rec(0, 0, 0x0FA0, "슬라이드 본문".encode("utf-16-le"))
                + rec(0, 0, 0x0FA8, b"ASCII text"))
    out: list[str] = []
    ex._parse_ppt_records(master + slide, out)
    assert out == ["슬라이드 본문", "ASCII text"]


def _soffice_convert(src: Path, fmt: str) -> Path:
    subprocess.run(["soffice", "--headless", "--convert-to", fmt, "--outdir", str(src.parent), str(src)],
                   check=True, capture_output=True, timeout=240)
    return src.with_suffix("." + fmt)


@pytest.mark.skipif(not HAS_SOFFICE, reason="LibreOffice 없음")
def test_legacy_xls_doc_ppt_with_real_files(tmp_path):
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation

    wb = Workbook()
    wb.active.append(["항목", "값"])
    wb.active.append(["가나다", 7])
    wb.save(tmp_path / "t.xlsx")
    xls = _soffice_convert(tmp_path / "t.xlsx", "xls")
    assert ex.extract_text(xls) == "항목, 값\n가나다, 7"

    d = Document()
    d.add_paragraph("레거시 워드 문서입니다. Second sentence here.")
    d.add_paragraph("둘째 문단")
    d.save(tmp_path / "t.docx")
    doc = _soffice_convert(tmp_path / "t.docx", "doc")
    builtin = ex.clean_text(ex._from_doc_builtin(doc))
    assert builtin == "레거시 워드 문서입니다. Second sentence here.\n둘째 문단"
    assert ex.extract_text(doc) == builtin

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "옛날 파워포인트"
    s.placeholders[1].text = "슬라이드 내용 하나"
    prs.save(tmp_path / "t.pptx")
    ppt = _soffice_convert(tmp_path / "t.pptx", "ppt")
    out = ex.clean_text(ex._from_ppt_builtin(ppt))
    assert "옛날 파워포인트" in out and "슬라이드 내용 하나" in out
    assert "Click to edit" not in out
    assert "옛날 파워포인트" in ex.extract_text(ppt)


def test_unsupported_and_missing(tmp_path):
    with pytest.raises(ex.ExtractError):
        ex.extract_text(tmp_path / "nope.pdf")
    p = tmp_path / "x.abc"
    p.write_text("x")
    with pytest.raises(ex.ExtractError):
        ex.extract_text(p)
